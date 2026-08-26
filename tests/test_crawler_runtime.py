from __future__ import annotations

import sqlite3
import sys
import tempfile
import types
import unittest
from contextlib import redirect_stdout
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch


try:
    import bs4  # noqa: F401
except ImportError:
    bs4_module = types.ModuleType("bs4")
    element_module = types.ModuleType("bs4.element")
    bs4_module.BeautifulSoup = object
    element_module.Tag = object
    sys.modules["bs4"] = bs4_module
    sys.modules["bs4.element"] = element_module


try:
    import pymupdf  # noqa: F401
except ImportError:
    pymupdf_module = types.ModuleType("pymupdf")

    class FakeTools:
        @staticmethod
        def mupdf_display_errors(_enabled: bool) -> None:
            return None

        @staticmethod
        def mupdf_display_warnings(_enabled: bool) -> None:
            return None

    pymupdf_module.TOOLS = FakeTools()
    pymupdf_module.open = lambda **_kwargs: None
    sys.modules["pymupdf"] = pymupdf_module


try:
    import requests  # noqa: F401
except ImportError:
    # Заглушки кладутся в sys.modules, а он один на процесс: их получает
    # не только этот модуль, но и всё, что импортируется после. Значит
    # заглушка обязана вести себя как класс, а не как метка.
    #
    # Раньше здесь стояло `Session = object`, и на машине без requests
    # это ломало совсем другой тест: `app_v2/es_http.py` при импорте
    # создаёт клиента и вызывает `Retry(total=2, connect=2, ...)`, а
    # `object()` аргументов не принимает. Падал `test_document_titles`,
    # с сообщением «object() takes no arguments» и трассировкой, в
    # которой не было ни слова про заглушки.
    #
    # Замечено только тогда, когда набор впервые запустили там, где
    # библиотек действительно нет.
    class StubObject:
        """Принимает любые аргументы и ничего не делает."""

        def __init__(self, *args, **kwargs) -> None:
            self.args = args
            self.kwargs = kwargs

        def __call__(self, *args, **kwargs):
            return self

        def __getattr__(self, name):
            return StubObject()

    requests_module = types.ModuleType("requests")
    requests_adapters_module = types.ModuleType(
        "requests.adapters"
    )
    urllib3_module = types.ModuleType("urllib3")
    urllib3_util_module = types.ModuleType("urllib3.util")
    urllib3_retry_module = types.ModuleType("urllib3.util.retry")

    requests_module.Session = StubObject
    requests_adapters_module.HTTPAdapter = StubObject
    urllib3_retry_module.Retry = StubObject

    sys.modules["requests"] = requests_module
    sys.modules["requests.adapters"] = requests_adapters_module
    sys.modules["urllib3"] = urllib3_module
    sys.modules["urllib3.util"] = urllib3_util_module
    sys.modules["urllib3.util.retry"] = urllib3_retry_module


from crawler_v2.content_processing import (  # noqa: E402
    CorruptPdfError,
    ExtractedContent,
    PDF_OCR_MODE_BOUNDED,
    PDF_OCR_MODE_FAST,
    _extract_pdf_in_process,
    extract_content,
    extract_pdf,
)
from crawler_v2.incremental_worker import (  # noqa: E402
    apply_authoritative_attachment_metadata,
    best_content_title,
    local_worker_is_alive,
    process_row,
    release_orphaned_local_locks,
    resolve_attachment_parent_metadata,
    useful_pdf_title,
)
from crawler_v2.incremental_db import (  # noqa: E402
    IncrementalStateDatabase,
)
from crawler_v2.legal_metadata import (  # noqa: E402
    extract_legal_metadata,
    extract_legal_title_metadata,
    parse_russian_date,
)
from crawler_v2.pdf_ocr import (  # noqa: E402
    PDF_EXTRACTION_VERSION,
    PDF_FAST_EXTRACTION_VERSION,
    PdfPageText,
    evaluate_text_quality,
    run_ocr_self_test,
)


CRAWL_SCHEMA = """
CREATE TABLE crawl_urls (
    url TEXT PRIMARY KEY,
    status TEXT NOT NULL,
    is_active INTEGER NOT NULL,
    doc_type TEXT,
    next_check_at TEXT,
    last_checked_at TEXT,
    locked_at TEXT,
    locked_by TEXT,
    last_error TEXT,
    http_status INTEGER,
    etag TEXT,
    last_modified TEXT,
    content_hash TEXT,
    content_type TEXT,
    last_fetched_at TEXT,
    last_indexed_at TEXT,
    last_seen_at TEXT,
    attempt_count INTEGER NOT NULL DEFAULT 0,
    failure_count INTEGER NOT NULL DEFAULT 0,
    indexed_chunks INTEGER NOT NULL DEFAULT 0,
    consecutive_missing INTEGER NOT NULL DEFAULT 0,
    extraction_version TEXT,
    page_url TEXT,
    file_url TEXT,
    parent_url TEXT,
    document_number TEXT,
    document_date TEXT,
    document_kind TEXT
);
"""


class FakePage:
    def __init__(self, text: str) -> None:
        self.text = text

    def get_text(self, _kind: str, *, sort: bool) -> str:
        self.sort = sort
        return self.text


class FakePdf:
    metadata = {"title": "Тестовый PDF"}
    page_count = 2

    def __init__(self, *, all_broken: bool = False) -> None:
        self.all_broken = all_broken
        self.closed = False

    def load_page(self, index: int) -> FakePage:
        if self.all_broken or index == 0:
            raise RuntimeError("non-page object in page tree")

        return FakePage(
            "Исправная страница документа " * 10
        )

    def close(self) -> None:
        self.closed = True


class FakeReferencePdf:
    metadata = {"title": "Информационный материал"}
    page_count = 2

    def __init__(self) -> None:
        self.closed = False

    def load_page(self, index: int) -> FakePage:
        if index == 0:
            return FakePage(
                "Информационный материал для сотрудников города Сочи. "
                "Документ содержит справочные сведения и рекомендации. "
                "Основной текст первой страницы не является правовым актом. "
            )

        return FakePage(
            "В работе следует учитывать Указ Президента Российской "
            "Федерации от 29 декабря 2022 года № 968. "
            "Эта ссылка приведена только как справочная информация. "
        )

    def close(self) -> None:
        self.closed = True


class FakeOcrPage:
    rotation = 0

    def __init__(self) -> None:
        self.ocr_calls = 0

    def set_rotation(self, value: int) -> None:
        self.rotation = value

    def get_textpage_ocr(self, **_kwargs: object) -> str:
        self.ocr_calls += 1
        return "ocr-text-page"

    def get_text(
        self,
        _kind: str,
        *,
        sort: bool,
        textpage: object = None,
    ) -> str:
        self.sort = sort

        if textpage is not None:
            return (
                "АДМИНИСТРАЦИЯ ГОРОДА СОЧИ ПОСТАНОВЛЕНИЕ "
                "от 12 февраля 2019 года № 223. "
                "О внесении изменений в постановление администрации "
                "города Сочи от 31 марта 2016 года № 788. "
            ) * 2

        return (
            "AAMrHrrCTPAqU TOPOAA COrrU TIO "
            "CTAIIOBJIEIII,IE sHeceuuu H3MeHeHrIfi "
        ) * 4


class FakeOcrPdf:
    metadata = {
        "title": (
            "Официальное опубликование "
            "муниципальных правовых актов"
        )
    }
    page_count = 1

    def __init__(self) -> None:
        self.page = FakeOcrPage()
        self.closed = False

    def load_page(self, _index: int) -> FakeOcrPage:
        return self.page

    def close(self) -> None:
        self.closed = True


class FakePartialHeaderPage:
    rotation = 0

    def set_rotation(self, value: int) -> None:
        self.rotation = value

    def get_textpage_ocr(self, **_kwargs: object) -> str:
        return "ocr-header-page"

    def get_text(
        self,
        _kind: str,
        *,
        sort: bool,
        textpage: object = None,
    ) -> str:
        self.sort = sort

        if textpage is not None:
            return (
                "АДМИНИСТРАЦИЯ ГОРОДА СОЧИ РАСПОРЯЖЕНИЕ "
                "от 30.11.2012 № 620-р город Сочи. "
                "О проведении субботника 1 и 2 декабря 2012 года "
                "на территории города Сочи. В целях обеспечения "
                "санитарно-эпидемиологического благополучия. "
            ) * 2

        return (
            "АДМИНИСТРАЦИЯ ГОРОДА СОЧИ РАСПОРЯЖЕНИЕ "
            "от за //.ж/л № ejo-p город Сочи. "
            "О проведении субботника 1 и 2 декабря 2012 года "
            "на территории города Сочи. В целях обеспечения "
            "санитарно-эпидемиологического благополучия. "
        ) * 3


class FakePartialHeaderPdf:
    metadata = {
        "title": "3de4a95792b3857abbaa210d20d655d3"
    }
    page_count = 1

    def __init__(self) -> None:
        self.page = FakePartialHeaderPage()
        self.closed = False

    def load_page(self, _index: int) -> FakePartialHeaderPage:
        return self.page

    def close(self) -> None:
        self.closed = True


class FakeRejectedHeaderPage:
    rotation = 0

    def set_rotation(self, value: int) -> None:
        self.rotation = value

    def get_textpage_ocr(self, **_kwargs: object) -> str:
        return "ocr-rejected-header-page"

    def get_text(
        self,
        _kind: str,
        *,
        sort: bool,
        textpage: object = None,
    ) -> str:
        self.sort = sort

        if textpage is not None:
            return (
                "АДМИНИСТРАЦИЯ ГОРОДА СОЧИ РАСПОРЯЖЕНИЕ "
                "от За №6209 город Сочи. О проведении субботника "
                "1 и 2 декабря 2012 года на территории города Сочи. "
                "В целях обеспечения санитарно-эпидемиологического "
                "благополучия населения и отдыхающих. "
            ) * 4

        return (
            "2 организовать работу по проведению субботников. "
            "АДМИНИСТРАЦИЯ ГОРОДА СОЧИ РАСПОРЯЖЕНИЕ "
            "от за //.ж/л № ejo-p город Сочи. О проведении "
            "субботника 1 и 2 декабря 2012 года на территории "
            "города Сочи. В целях обеспечения санитарно-"
            "эпидемиологического благополучия населения. "
        ) * 4


class FakeRejectedHeaderPdf:
    metadata = {
        "title": "3de4a95792b3857abbaa210d20d655d3"
    }
    page_count = 1

    def __init__(self) -> None:
        self.page = FakeRejectedHeaderPage()
        self.closed = False

    def load_page(self, _index: int) -> FakeRejectedHeaderPage:
        return self.page

    def close(self) -> None:
        self.closed = True


class CrawlerRuntimeTests(unittest.TestCase):
    @staticmethod
    def parent_card_text() -> str:
        return """
        Номер документа:
        Дата документа:
        30 ноября 2012
        Опубликовано на сайте: 26 марта 2013
        Тип документа:
        Скачать документ Формат файла .pdf Размер 1.05 МБ
        Распоряжение администрации города No.620-р от 30.11.2012
        "О проведении субботника 1 и 2 декабря 2012 года на
        территории города Сочи"
        """

    def test_parent_card_accepts_no_dot_document_number_marker(self) -> None:
        container = types.SimpleNamespace(
            select=lambda _selector: ()
        )
        metadata = extract_legal_metadata(
            types.SimpleNamespace(),
            container,
            text=self.parent_card_text(),
            final_url=(
                "https://sochi.ru/gorodskaya-vlast/"
                "normativno-pravovyye-akty/?ELEMENT_ID=4638"
            ),
        )

        self.assertEqual(metadata.document_number, "620-р")
        self.assertEqual(metadata.document_date, "2012-11-30")
        self.assertEqual(metadata.document_kind, "Распоряжение")

    def test_impossible_calendar_dates_are_rejected(self) -> None:
        self.assertIsNone(
            parse_russian_date(
                "31.02.2012"
            )
        )
        self.assertIsNone(
            parse_russian_date(
                "32 декабря 2012"
            )
        )
        self.assertEqual(
            parse_russian_date(
                "30.11.2012"
            ),
            "2012-11-30",
        )

    def test_attachment_title_extracts_ukaz_metadata(self) -> None:
        metadata = extract_legal_title_metadata(
            "Указ Президента Российской Федерации от 29 декабря "
            "2022 года № 968 «Об особенностях исполнения обязанностей»"
        )

        self.assertEqual(metadata.document_number, "968")
        self.assertEqual(metadata.document_date, "2022-12-29")
        self.assertEqual(metadata.document_kind, "Указ")

    def test_generic_attachment_title_clears_false_pdf_metadata(self) -> None:
        content = ExtractedContent(
            title="Перечень документов",
            section="PDF-документ",
            doc_type="pdf",
            content_type="application/pdf",
            pages=((1, "Справочный текст"),),
            document_number="870",
            document_date="2024-10-10",
            document_kind=None,
            extraction_version=PDF_EXTRACTION_VERSION,
        )
        row = {
            "url": "https://sochi.ru/upload/list.pdf",
            "parent_url": "https://sochi.ru/municipal-service/",
            "attachment_title": (
                "80.24 КБ Перечень документов, необходимых для "
                "поступления на муниципальную службу"
            ),
            "_parent_document_number": None,
            "_parent_document_date": None,
            "_parent_document_kind": None,
        }

        merged = apply_authoritative_attachment_metadata(
            row,
            content,
        )

        self.assertIsNone(merged.document_number)
        self.assertIsNone(merged.document_date)
        self.assertIsNone(merged.document_kind)

    def test_ukaz_attachment_title_overrides_reference_inside_pdf(self) -> None:
        content = ExtractedContent(
            title="ukaz968.pdf",
            section="PDF-документ",
            doc_type="pdf",
            content_type="application/pdf",
            pages=((1, "Текст с упоминанием другого указа"),),
            document_number="557",
            document_date="2009-05-18",
            document_kind=None,
            extraction_version=PDF_EXTRACTION_VERSION,
        )
        row = {
            "url": "https://sochi.ru/ukaz968.pdf",
            "parent_url": "https://sochi.ru/anti-corruption/",
            "attachment_title": (
                "Указ Президента Российской Федерации от 29 декабря "
                "2022 года № 968 «Об особенностях исполнения обязанностей»"
            ),
            "_parent_document_number": None,
            "_parent_document_date": None,
            "_parent_document_kind": None,
        }

        merged = apply_authoritative_attachment_metadata(
            row,
            content,
        )

        self.assertEqual(merged.document_number, "968")
        self.assertEqual(merged.document_date, "2022-12-29")
        self.assertEqual(merged.document_kind, "Указ")

    def test_new_html_then_pdf_precede_retry_and_missing(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "state.sqlite3"

            with sqlite3.connect(path) as connection:
                connection.executescript(CRAWL_SCHEMA)
                connection.executemany(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type, next_check_at
                    )
                    VALUES (?, ?, 1, ?, 'ready')
                    """,
                    (
                        ("https://sochi.ru/retry", "retry", "page"),
                        ("https://sochi.ru/missing", "missing", "page"),
                        ("https://sochi.ru/new.pdf", "discovered", "pdf"),
                        ("https://sochi.ru/new/", "discovered", "page"),
                        (
                            "https://sochi.ru/existing/",
                            "indexed_existing",
                            "page",
                        ),
                    ),
                )

            rows = IncrementalStateDatabase(path).peek_candidates(
                limit=5,
                specific_url=None,
                force=True,
            )

            self.assertEqual(
                [row["url"] for row in rows],
                [
                    "https://sochi.ru/new/",
                    "https://sochi.ru/new.pdf",
                    "https://sochi.ru/existing/",
                    "https://sochi.ru/retry",
                    "https://sochi.ru/missing",
                ],
            )

    def test_claiming_one_exact_url_does_not_lock_the_batch(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "state.sqlite3"

            with sqlite3.connect(path) as connection:
                connection.executescript(CRAWL_SCHEMA)
                connection.executemany(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type, next_check_at
                    )
                    VALUES (?, 'discovered', 1, 'pdf', 'ready')
                    """,
                    (
                        ("https://sochi.ru/a.pdf",),
                        ("https://sochi.ru/b.pdf",),
                        ("https://sochi.ru/c.pdf",),
                    ),
                )

            database = IncrementalStateDatabase(path)
            claimed = database.claim_candidates(
                limit=1,
                worker_id="host:123",
                specific_url="https://sochi.ru/b.pdf",
                exact_url=True,
                force=True,
            )

            self.assertEqual(
                [row["url"] for row in claimed],
                ["https://sochi.ru/b.pdf"],
            )

            with sqlite3.connect(path) as connection:
                states = dict(
                    connection.execute(
                        "SELECT url, status FROM crawl_urls"
                    ).fetchall()
                )

            self.assertEqual(states["https://sochi.ru/b.pdf"], "processing")
            self.assertEqual(states["https://sochi.ru/a.pdf"], "discovered")
            self.assertEqual(states["https://sochi.ru/c.pdf"], "discovered")

    def test_orphaned_local_worker_locks_are_released_immediately(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "state.sqlite3"
            proc_root = root / "proc"
            live_process = proc_root / "200"
            live_process.mkdir(parents=True)
            (live_process / "cmdline").write_bytes(
                b"python\0-m\0crawler_v2.incremental_worker\0"
            )

            with sqlite3.connect(path) as connection:
                connection.executescript(CRAWL_SCHEMA)
                connection.executemany(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type,
                        next_check_at, locked_at, locked_by
                    )
                    VALUES (?, 'processing', 1, 'pdf', 'later', 'now', ?)
                    """,
                    (
                        ("https://sochi.ru/dead.pdf", "host:100"),
                        ("https://sochi.ru/live.pdf", "host:200"),
                        ("https://sochi.ru/remote.pdf", "remote:300"),
                    ),
                )

            database = IncrementalStateDatabase(path)
            released = release_orphaned_local_locks(
                database,
                hostname="host",
                proc_root=proc_root,
            )

            self.assertEqual(released, 1)
            self.assertTrue(
                local_worker_is_alive(
                    "host:200",
                    hostname="host",
                    proc_root=proc_root,
                )
            )
            self.assertFalse(
                local_worker_is_alive(
                    "host:100",
                    hostname="host",
                    proc_root=proc_root,
                )
            )
            self.assertIsNone(
                local_worker_is_alive(
                    "remote:300",
                    hostname="host",
                    proc_root=proc_root,
                )
            )

            with sqlite3.connect(path) as connection:
                rows = {
                    row[0]: row[1:]
                    for row in connection.execute(
                        """
                        SELECT url, status, locked_by
                        FROM crawl_urls
                        """
                    ).fetchall()
                }

            self.assertEqual(
                rows["https://sochi.ru/dead.pdf"],
                ("retry", None),
            )
            self.assertEqual(
                rows["https://sochi.ru/live.pdf"],
                ("processing", "host:200"),
            )
            self.assertEqual(
                rows["https://sochi.ru/remote.pdf"],
                ("processing", "remote:300"),
            )

            recovered = database.release_all_processing(
                reason="installer recovery",
            )
            self.assertEqual(recovered, 2)

            with sqlite3.connect(path) as connection:
                processing = connection.execute(
                    """
                    SELECT COUNT(*)
                    FROM crawl_urls
                    WHERE status = 'processing'
                    """
                ).fetchone()[0]

            self.assertEqual(processing, 0)

    def test_pdf_candidate_contains_current_parent_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "state.sqlite3"
            parent_url = "https://sochi.ru/legal/?ELEMENT_ID=4638"
            pdf_url = "https://sochi.ru/upload/document.pdf"

            with sqlite3.connect(path) as connection:
                connection.executescript(CRAWL_SCHEMA)
                connection.executemany(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type, next_check_at,
                        parent_url, document_number, document_date,
                        document_kind
                    )
                    VALUES (?, 'indexed', 1, ?, 'ready', ?, ?, ?, ?)
                    """,
                    (
                        (
                            parent_url,
                            "page",
                            None,
                            "620-р",
                            "2012-11-30",
                            "Распоряжение",
                        ),
                        (
                            pdf_url,
                            "pdf",
                            parent_url,
                            "122",
                            "2012-12-02",
                            "Распоряжение",
                        ),
                    ),
                )

            rows = IncrementalStateDatabase(path).peek_candidates(
                limit=1,
                specific_url=pdf_url,
                force=True,
            )

            self.assertEqual(len(rows), 1)
            self.assertEqual(
                rows[0]["_parent_document_number"],
                "620-р",
            )
            self.assertEqual(
                rows[0]["_parent_document_date"],
                "2012-11-30",
            )
            self.assertEqual(
                rows[0]["_parent_document_kind"],
                "Распоряжение",
            )

    def test_new_html_pdf_and_files_are_interleaved(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "state.sqlite3"

            with sqlite3.connect(path) as connection:
                connection.executescript(CRAWL_SCHEMA)
                connection.executemany(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type, next_check_at
                    )
                    VALUES (?, 'discovered', 1, ?, 'ready')
                    """,
                    (
                        ("https://sochi.ru/page-a/", "page"),
                        ("https://sochi.ru/page-b/", "page"),
                        ("https://sochi.ru/page-c/", "page"),
                        ("https://sochi.ru/a.pdf", "pdf"),
                        ("https://sochi.ru/b.pdf", "pdf"),
                        ("https://sochi.ru/a.docx", "file"),
                        ("https://sochi.ru/b.docx", "file"),
                    ),
                )

            rows = IncrementalStateDatabase(path).peek_candidates(
                limit=6,
                specific_url=None,
                force=True,
            )

            self.assertEqual(
                [row["doc_type"] for row in rows],
                ["page", "pdf", "file", "page", "pdf", "file"],
            )

    def test_pdf_ocr_backfill_selects_each_active_pdf_once(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "state.sqlite3"

            with sqlite3.connect(path) as connection:
                connection.executescript(CRAWL_SCHEMA)
                connection.executemany(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type,
                        next_check_at, extraction_version
                    )
                    VALUES (?, 'indexed', 1, ?, 'later', ?)
                    """,
                    (
                        (
                            "https://sochi.ru/old.pdf",
                            "pdf",
                            None,
                        ),
                        (
                            "https://sochi.ru/current.pdf",
                            "pdf",
                            PDF_EXTRACTION_VERSION,
                        ),
                        (
                            "https://sochi.ru/page/",
                            "page",
                            None,
                        ),
                        (
                            "https://sochi.ru/not-indexed.pdf",
                            "pdf",
                            None,
                        ),
                    ),
                )

                connection.execute(
                    """
                    UPDATE crawl_urls
                    SET status = 'discovered'
                    WHERE url = 'https://sochi.ru/not-indexed.pdf'
                    """
                )

            rows = IncrementalStateDatabase(
                path
            ).peek_candidates(
                limit=10,
                specific_url=None,
                force=True,
                required_extraction_version=(
                    PDF_EXTRACTION_VERSION
                ),
                require_existing_index=True,
            )

            self.assertEqual(
                [row["url"] for row in rows],
                ["https://sochi.ru/old.pdf"],
            )

    def test_unsupported_and_corrupt_are_terminal(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "state.sqlite3"

            with sqlite3.connect(path) as connection:
                connection.executescript(CRAWL_SCHEMA)
                connection.execute(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type, next_check_at
                    )
                    VALUES ('https://sochi.ru/bad.pdf', 'processing', 1, 'pdf', 'later')
                    """
                )

            database = IncrementalStateDatabase(path)
            database.mark_skipped(
                "https://sochi.ru/bad.pdf",
                status="skipped_corrupt",
                checked_at="now",
                next_check_at="later",
                error_message="broken",
            )

            with sqlite3.connect(path) as connection:
                state = connection.execute(
                    """
                    SELECT status, is_active, next_check_at
                    FROM crawl_urls
                    """
                ).fetchone()

            self.assertEqual(state, ("skipped_corrupt", 0, None))

    def test_source_hash_is_persisted_without_replacing_ocr_state(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "state.sqlite3"

            with sqlite3.connect(path) as connection:
                connection.executescript(CRAWL_SCHEMA)
                connection.execute(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type,
                        extraction_version, content_hash,
                        indexed_chunks
                    )
                    VALUES (?, 'processing', 1, 'pdf', ?, 'ocr-hash', 7)
                    """,
                    (
                        "https://sochi.ru/source.pdf",
                        PDF_EXTRACTION_VERSION,
                    ),
                )

            database = IncrementalStateDatabase(path)
            database.migrate()
            database.mark_source_unchanged(
                "https://sochi.ru/source.pdf",
                checked_at="checked",
                next_check_at="next",
                fetched_at="fetched",
                source_hash="source-hash",
                etag="etag",
                last_modified="modified",
            )

            with sqlite3.connect(path) as connection:
                row = connection.execute(
                    """
                    SELECT
                        status,
                        extraction_version,
                        content_hash,
                        indexed_chunks,
                        source_hash,
                        locked_at,
                        locked_by
                    FROM crawl_urls
                    """
                ).fetchone()

            self.assertEqual(
                row,
                (
                    "unchanged",
                    PDF_EXTRACTION_VERSION,
                    "ocr-hash",
                    7,
                    "source-hash",
                    None,
                    None,
                ),
            )

            with sqlite3.connect(path) as connection:
                connection.execute(
                    """
                    INSERT INTO crawl_urls (
                        url, status, is_active, doc_type
                    )
                    VALUES (?, 'processing', 1, 'pdf')
                    """,
                    ("https://sochi.ru/new.pdf",),
                )

            database.mark_indexed(
                "https://sochi.ru/new.pdf",
                checked_at="checked",
                next_check_at="next",
                fetched_at="fetched",
                indexed_at="indexed",
                content_hash="content-hash",
                content_type="application/pdf",
                doc_type="pdf",
                indexed_chunks=3,
                etag=None,
                last_modified=None,
                final_url="https://sochi.ru/new.pdf",
                extraction_version=(
                    PDF_FAST_EXTRACTION_VERSION
                ),
                source_hash="new-source-hash",
            )

            with sqlite3.connect(path) as connection:
                indexed = connection.execute(
                    """
                    SELECT
                        status,
                        extraction_version,
                        source_hash,
                        indexed_chunks
                    FROM crawl_urls
                    WHERE url = 'https://sochi.ru/new.pdf'
                    """
                ).fetchone()

            self.assertEqual(
                indexed,
                (
                    "indexed",
                    PDF_FAST_EXTRACTION_VERSION,
                    "new-source-hash",
                    3,
                ),
            )

    def test_pdf_salvages_good_pages_and_terminally_classifies_bad_tree(self) -> None:
        partially_broken = FakePdf()

        with patch(
            "crawler_v2.content_processing.pymupdf.open",
            return_value=partially_broken,
        ):
            content = extract_pdf(
                b"%PDF-test",
                content_type="application/pdf",
                final_url="https://sochi.ru/test.pdf",
            )

        self.assertEqual(len(content.pages), 1)
        self.assertEqual(content.pages[0][0], 2)
        self.assertTrue(partially_broken.closed)

        completely_broken = FakePdf(all_broken=True)

        with patch(
            "crawler_v2.content_processing.pymupdf.open",
            return_value=completely_broken,
        ):
            with self.assertRaises(CorruptPdfError):
                extract_pdf(
                    b"%PDF-test",
                    content_type="application/pdf",
                    final_url="https://sochi.ru/test.pdf",
                )

        self.assertTrue(completely_broken.closed)

    def test_pdf_does_not_take_metadata_from_later_reference_page(self) -> None:
        reference_pdf = FakeReferencePdf()

        with patch(
            "crawler_v2.content_processing.pymupdf.open",
            return_value=reference_pdf,
        ):
            content = extract_pdf(
                b"%PDF-reference",
                content_type="application/pdf",
                final_url="https://sochi.ru/reference.pdf",
            )

        self.assertIsNone(content.document_number)
        self.assertIsNone(content.document_date)
        self.assertIsNone(content.document_kind)
        self.assertTrue(reference_pdf.closed)

    @unittest.skip(
        "движок OCR заменён: страница рендерится в растр и уходит во внешний tesseract. Заглушки ниже описывают снятый путь через page.get_textpage_ocr. Замена — tests/test_v250_units.py, классы с pdf_ocr_engine и pdf_ocr_image"
    )
    def test_pdf_replaces_long_broken_glyph_layer_with_russian_ocr(self) -> None:
        fake_pdf = FakeOcrPdf()

        with patch(
            "crawler_v2.content_processing.pymupdf.open",
            return_value=fake_pdf,
        ):
            content = extract_pdf(
                b"%PDF-test",
                content_type="application/pdf",
                final_url="https://sochi.ru/test.pdf",
            )

        self.assertEqual(
            content.extraction_version,
            PDF_EXTRACTION_VERSION,
        )
        self.assertEqual(content.ocr_pages, (1,))
        self.assertEqual(content.unreadable_pages, ())
        self.assertEqual(
            content.document_number,
            "223",
        )
        self.assertEqual(
            content.document_date,
            "2019-02-12",
        )
        self.assertEqual(
            content.document_kind,
            "Постановление",
        )
        self.assertIn(
            "АДМИНИСТРАЦИЯ ГОРОДА СОЧИ",
            content.pages[0][1],
        )
        self.assertNotIn(
            "AAMrHrrCTPAqU",
            content.pages[0][1],
        )
        self.assertTrue(fake_pdf.closed)

    def test_fast_pdf_mode_defers_ocr_and_returns_immediately(self) -> None:
        fake_pdf = FakeOcrPdf()

        with patch(
            "crawler_v2.content_processing.pymupdf.open",
            return_value=fake_pdf,
        ):
            content = extract_pdf(
                b"%PDF-fast",
                content_type="application/pdf",
                final_url="https://sochi.ru/fast.pdf",
                ocr_mode=PDF_OCR_MODE_FAST,
            )

        self.assertEqual(
            content.extraction_version,
            PDF_FAST_EXTRACTION_VERSION,
        )
        self.assertEqual(content.ocr_pages, ())
        self.assertEqual(content.ocr_attempted_pages, ())
        self.assertEqual(content.ocr_deferred_pages, (1,))
        self.assertEqual(content.unreadable_pages, (1,))
        self.assertEqual(fake_pdf.page.ocr_calls, 0)
        self.assertTrue(fake_pdf.closed)

    def test_fast_recheck_preserves_completed_ocr_when_source_is_same(self) -> None:
        class RecordingDatabase:
            def __init__(self) -> None:
                self.calls: list[dict[str, object]] = []

            def mark_source_unchanged(
                self,
                _url: str,
                **kwargs: object,
            ) -> None:
                self.calls.append(kwargs)

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.pdf"
            source = b"%PDF-unchanged-source"
            path.write_bytes(source)

            import hashlib

            source_hash = hashlib.sha256(
                source
            ).hexdigest()
            result = types.SimpleNamespace(
                status_code=200,
                final_url="https://sochi.ru/source.pdf",
                content_type="application/pdf",
                body=b"",
                payload=path,
                size_bytes=len(source),
                etag="etag",
                last_modified="date",
                cleanup=lambda: path.unlink(
                    missing_ok=True
                ),
            )
            database = RecordingDatabase()
            row = {
                "url": result.final_url,
                "file_url": result.final_url,
                "status": "indexed",
                "source_hash": source_hash,
                "extraction_version": (
                    PDF_EXTRACTION_VERSION
                ),
            }

            with patch(
                "crawler_v2.incremental_worker.extract_content"
            ) as extract_mock, redirect_stdout(StringIO()):
                status = process_row(
                    row,
                    database=database,
                    legal_database=types.SimpleNamespace(),
                    fetcher=types.SimpleNamespace(
                        fetch=lambda *_args, **_kwargs: result
                    ),
                    writer=types.SimpleNamespace(),
                    dry_run=False,
                    ignore_http_cache=True,
                    pdf_ocr_mode=PDF_OCR_MODE_FAST,
                )

            self.assertEqual(status, "unchanged")
            self.assertEqual(len(database.calls), 1)
            self.assertEqual(
                database.calls[0]["source_hash"],
                source_hash,
            )
            extract_mock.assert_not_called()
            self.assertFalse(path.exists())

    @unittest.skip(
        "движок OCR заменён: страница рендерится в растр и уходит во внешний tesseract. Заглушки ниже описывают снятый путь через page.get_textpage_ocr. Замена — tests/test_v250_units.py, классы с pdf_ocr_engine и pdf_ocr_image"
    )
    def test_page_timeout_keeps_other_pdf_pages(self) -> None:
        class MixedPdf:
            metadata = {"title": "Смешанный PDF"}
            page_count = 2

            def __init__(self) -> None:
                self.closed = False

            @staticmethod
            def load_page(index: int):
                if index == 0:
                    return FakeOcrPage()

                return FakePage(
                    "Вторая читаемая страница документа города Сочи. "
                    * 8
                )

            def close(self) -> None:
                self.closed = True

        mixed_pdf = MixedPdf()
        broken_text = (
            "AAMrHrrCTPAqU TOPOAA COrrU "
            * 8
        )
        timeout_result = PdfPageText(
            text="",
            method="unreadable",
            quality=evaluate_text_quality(
                broken_text
            ),
            ocr_attempted=True,
            error=(
                "timeout: OCR страницы превысил 30 секунд"
            ),
        )

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mixed.pdf"
            path.write_bytes(b"%PDF-mixed")

            with patch(
                "crawler_v2.content_processing.pymupdf.open",
                return_value=mixed_pdf,
            ), patch(
                "crawler_v2.content_processing._extract_pdf_page_isolated",
                return_value=timeout_result,
            ):
                content = _extract_pdf_in_process(
                    path,
                    content_type="application/pdf",
                    final_url="https://sochi.ru/mixed.pdf",
                    ocr_mode=PDF_OCR_MODE_BOUNDED,
                )

        self.assertEqual(content.ocr_timed_out_pages, (1,))
        self.assertEqual(content.ocr_rejected_pages, (1,))
        self.assertEqual(content.unreadable_pages, (1,))
        self.assertEqual(
            [page_number for page_number, _text in content.pages],
            [2],
        )
        self.assertTrue(mixed_pdf.closed)

    @unittest.skip(
        "движок OCR заменён: страница рендерится в растр и уходит во внешний tesseract. Заглушки ниже описывают снятый путь через page.get_textpage_ocr. Замена — tests/test_v250_units.py, классы с pdf_ocr_engine и pdf_ocr_image"
    )
    def test_pdf_ocr_repairs_damaged_header_on_otherwise_readable_page(self) -> None:
        fake_pdf = FakePartialHeaderPdf()

        with patch(
            "crawler_v2.content_processing.pymupdf.open",
            return_value=fake_pdf,
        ):
            content = extract_pdf(
                b"%PDF-partial-header",
                content_type="application/pdf",
                final_url=(
                    "https://sochi.ru/upload/iblock/3de/"
                    "3de4a95792b3857abbaa210d20d655d3"
                ),
            )

        self.assertEqual(content.ocr_pages, (1,))
        self.assertEqual(content.document_number, "620-р")
        self.assertEqual(content.document_date, "2012-11-30")
        self.assertEqual(content.document_kind, "Распоряжение")
        self.assertIn("№ 620-р", content.pages[0][1])
        self.assertNotIn("№ ejo-p", content.pages[0][1])
        self.assertTrue(fake_pdf.closed)

    @unittest.skip(
        "движок OCR заменён: страница рендерится в растр и уходит во внешний tesseract. Заглушки ниже описывают снятый путь через page.get_textpage_ocr. Замена — tests/test_v250_units.py, классы с pdf_ocr_engine и pdf_ocr_image"
    )
    def test_unreliable_realistic_ocr_header_is_reported_and_rejected(self) -> None:
        fake_pdf = FakeRejectedHeaderPdf()

        with patch(
            "crawler_v2.content_processing.pymupdf.open",
            return_value=fake_pdf,
        ):
            content = extract_pdf(
                b"%PDF-realistic-rejected-header",
                content_type="application/pdf",
                final_url=(
                    "https://sochi.ru/upload/iblock/3de/"
                    "3de4a95792b3857abbaa210d20d655d3"
                ),
            )

        self.assertEqual(content.ocr_pages, ())
        self.assertEqual(content.ocr_attempted_pages, (1,))
        self.assertEqual(content.ocr_rejected_pages, (1,))
        self.assertIn(
            "не восстановил надёжные номер и дату",
            content.ocr_rejection_reasons[0],
        )
        self.assertIn("№ ejo-p", content.pages[0][1])
        self.assertNotIn("№6209", content.pages[0][1])
        self.assertTrue(fake_pdf.closed)

    def test_incomplete_attachment_metadata_is_loaded_from_parent_card(self) -> None:
        parent_url = (
            "https://sochi.ru/gorodskaya-vlast/"
            "normativno-pravovyye-akty/?ELEMENT_ID=4638"
        )
        row = {
            "url": "https://sochi.ru/upload/document",
            "parent_url": parent_url,
            "document_number": None,
            "document_date": "2012-11-30",
            "document_kind": "Распоряжение",
            "_parent_document_number": None,
            "_parent_document_date": "2012-11-30",
            "_parent_document_kind": "Распоряжение",
        }
        content = ExtractedContent(
            title="3de4a95792b3857abbaa210d20d655d3",
            section="PDF-документ",
            doc_type="pdf",
            content_type="application/pdf",
            pages=((1, "Читаемый текст правового акта"),),
            document_number=None,
            document_date=None,
            document_kind="Распоряжение",
            extraction_version=PDF_EXTRACTION_VERSION,
            ocr_attempted_pages=(1,),
            ocr_rejected_pages=(1,),
        )
        parent_result = types.SimpleNamespace(
            status_code=200,
            content_type="text/html",
            final_url=parent_url,
            body=b"parent-card",
        )
        fetcher = types.SimpleNamespace(
            fetch=lambda *_args, **_kwargs: parent_result
        )
        parent_content = ExtractedContent(
            title=(
                "Распоряжение администрации города "
                "No.620-р от 30.11.2012"
            ),
            section="Правовые акты",
            doc_type="page",
            content_type="text/html",
            pages=((1, self.parent_card_text()),),
            document_number="620-р",
            document_date="2012-11-30",
            document_kind="Распоряжение",
        )

        with patch(
            "crawler_v2.incremental_worker.extract_content",
            return_value=parent_content,
        ), redirect_stdout(StringIO()):
            resolved, fetched = resolve_attachment_parent_metadata(
                row,
                content,
                fetcher=fetcher,
            )
        merged = apply_authoritative_attachment_metadata(
            resolved,
            content,
        )

        self.assertEqual(fetched["document_number"], "620-р")
        self.assertEqual(merged.document_number, "620-р")
        self.assertEqual(merged.document_date, "2012-11-30")
        self.assertEqual(merged.document_kind, "Распоряжение")

    def test_control_pdf_dry_run_reports_correct_parent_metadata_and_ocr_state(self) -> None:
        pdf_url = (
            "http://sochi.ru/upload/iblock/3de/"
            "3de4a95792b3857abbaa210d20d655d3"
        )
        parent_url = (
            "http://sochi.ru/gorodskaya-vlast/"
            "normativno-pravovyye-akty/?ELEMENT_ID=4638"
        )
        pdf_content = ExtractedContent(
            title="3de4a95792b3857abbaa210d20d655d3",
            section="PDF-документ",
            doc_type="pdf",
            content_type="application/pdf",
            pages=((1, "Читаемый текст распоряжения " * 20),),
            document_number=None,
            document_date=None,
            document_kind="Распоряжение",
            extraction_version=PDF_EXTRACTION_VERSION,
            ocr_pages=(),
            ocr_attempted_pages=(1,),
            ocr_rejected_pages=(1,),
            ocr_rejection_reasons=(
                "стр. 1: OCR шапки не восстановил надёжные номер и дату",
            ),
            source_page_count=1,
        )
        parent_content = ExtractedContent(
            title=(
                "Распоряжение администрации города "
                "No.620-р от 30.11.2012"
            ),
            section="Правовые акты",
            doc_type="page",
            content_type="text/html",
            pages=((1, self.parent_card_text()),),
            document_number="620-р",
            document_date="2012-11-30",
            document_kind="Распоряжение",
        )
        pdf_result = types.SimpleNamespace(
            status_code=200,
            final_url=pdf_url,
            content_type="application/pdf",
            body=b"%PDF-control",
            etag=None,
            last_modified=None,
        )
        parent_result = types.SimpleNamespace(
            status_code=200,
            final_url=parent_url,
            content_type="text/html",
            body=b"parent-card",
            etag=None,
            last_modified=None,
        )

        def fetch(url: str, **_kwargs: object) -> object:
            return parent_result if url == parent_url else pdf_result

        row = {
            "url": pdf_url,
            "file_url": pdf_url,
            "page_url": parent_url,
            "parent_url": parent_url,
            "doc_type": "pdf",
            "attachment_title": (
                "Скачать документ Формат файла .f1fee19f39809ebf61cebfc28d61e74f"
            ),
            "document_number": None,
            "document_date": "2012-11-30",
            "document_kind": "Распоряжение",
            "_parent_document_number": None,
            "_parent_document_date": "2012-11-30",
            "_parent_document_kind": "Распоряжение",
            "content_hash": "old-hash",
        }
        output = StringIO()

        with patch(
            "crawler_v2.incremental_worker.extract_content",
            side_effect=(pdf_content, parent_content),
        ), redirect_stdout(output):
            status = process_row(
                row,
                database=types.SimpleNamespace(),
                legal_database=types.SimpleNamespace(),
                fetcher=types.SimpleNamespace(fetch=fetch),
                writer=types.SimpleNamespace(),
                dry_run=True,
                ignore_http_cache=True,
            )

        rendered = output.getvalue()
        self.assertEqual(status, "would_index")
        self.assertIn("Номер документа: 620-р", rendered)
        self.assertIn("Дата документа: 2012-11-30", rendered)
        self.assertIn("PDF OCR страниц: 0", rendered)
        self.assertIn("PDF OCR попыток: 1", rendered)
        self.assertIn("PDF OCR отклонено: 1", rendered)
        self.assertIn(
            "Заголовок: Распоряжение № 620-р от 2012-11-30",
            rendered,
        )
        self.assertIn("[DRY RUN] Elasticsearch и SQLite не изменены", rendered)

    def test_parent_card_metadata_overrides_false_pdf_metadata(self) -> None:
        content = ExtractedContent(
            title="3de4a95792b3857abbaa210d20d655d3",
            section="PDF-документ",
            doc_type="pdf",
            content_type="application/pdf",
            pages=((1, "Читаемый текст правового акта"),),
            document_number="122",
            document_date="2012-12-02",
            document_kind="Распоряжение",
            extraction_version=PDF_EXTRACTION_VERSION,
        )
        row = {
            "url": "https://sochi.ru/upload/document",
            "parent_url": (
                "https://sochi.ru/gorodskaya-vlast/"
                "normativno-pravovyye-akty/?ELEMENT_ID=4638"
            ),
            "attachment_title": (
                "Официальное опубликование "
                "муниципальных правовых актов"
            ),
            "document_number": "620-р",
            "document_date": "2012-11-30",
            "document_kind": "Распоряжение",
            "_parent_document_number": "620-р",
            "_parent_document_date": "2012-11-30",
            "_parent_document_kind": "Распоряжение",
        }

        merged = apply_authoritative_attachment_metadata(
            row,
            content,
        )

        self.assertEqual(merged.document_number, "620-р")
        self.assertEqual(merged.document_date, "2012-11-30")
        self.assertEqual(merged.document_kind, "Распоряжение")
        self.assertEqual(
            best_content_title(row, merged),
            "Распоряжение № 620-р от 2012-11-30",
        )
        self.assertFalse(
            useful_pdf_title(
                "3de4a95792b3857abbaa210d20d655d3"
            )
        )

    def test_user_garbled_pdf_samples_are_rejected_by_quality_gate(self) -> None:
        samples = (
            (
                "ffi rlufl AAMI,IHIICTPA roPoAA coqu "
                "IIO CTAHOBJIEHIIE ffi 03 "
            ) * 4,
            (
                "soruoxBlI.H'v eVodor esurJ yvqecvrJVou "
                "ore rnV oc,ftrnc g reeur(rcs au orcEH "
            ) * 4,
            (
                "AAMrHrrCTPAqU TOPOAA COrrU TIO "
                "CTAIIOBJIEIII,IE sHeceuuu H3MeHeHrIfi "
            ) * 4,
        )

        for sample in samples:
            with self.subTest(sample=sample[:30]):
                quality = evaluate_text_quality(sample)
                self.assertFalse(quality.readable)

    @unittest.skip(
        "движок OCR заменён: страница рендерится в растр и уходит во внешний tesseract. Заглушки ниже описывают снятый путь через page.get_textpage_ocr. Замена — tests/test_v250_units.py, классы с pdf_ocr_engine и pdf_ocr_image"
    )
    def test_packaged_russian_english_ocr_runtime_is_operational(self) -> None:
        tessdata = (
            Path(__file__).resolve().parents[1]
            / "ocr"
            / "tessdata"
        )
        recognized = run_ocr_self_test(
            tessdata
        )

        self.assertIn(
            "OFFICIAL",
            recognized.upper(),
        )
        self.assertIn(
            "СОЧИ",
            recognized.upper(),
        )
        self.assertIn(
            "ПОСТАНОВЛЕНИЕ",
            recognized.upper(),
        )

    def test_docx_xlsx_and_cp1251_csv_are_extractable(self) -> None:
        try:
            from docx import Document
            from openpyxl import Workbook
        except ImportError:
            self.skipTest(
                "DOCX/XLSX libraries are checked by installer preflight"
            )

        docx_buffer = BytesIO()
        document = Document()
        document.add_paragraph(
            "Документ DOCX содержит полезный текст для внутреннего поиска. "
            * 3
        )
        document.save(docx_buffer)

        docx_content = extract_content(
            docx_buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            final_url="https://sochi.ru/files/test.docx",
        )

        self.assertEqual(docx_content.doc_type, "file")
        self.assertIn("полезный текст", docx_content.pages[0][1])

        xlsx_buffer = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Данные"
        sheet.append([
            "Строка таблицы для внутреннего поиска " * 3,
            42,
        ])
        workbook.save(xlsx_buffer)
        workbook.close()

        xlsx_content = extract_content(
            xlsx_buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            final_url="https://sochi.ru/files/test.xlsx",
        )

        self.assertEqual(xlsx_content.doc_type, "file")
        self.assertIn("Лист: Данные", xlsx_content.pages[0][1])

        csv_body = (
            "Название;Значение\n"
            + "Документ для внутреннего поиска;42\n" * 4
        ).encode("cp1251")
        csv_content = extract_content(
            csv_body,
            content_type="text/csv",
            final_url="https://sochi.ru/files/test.csv",
        )

        self.assertEqual(csv_content.doc_type, "file")
        self.assertIn("Документ для внутреннего поиска", csv_content.pages[0][1])


if __name__ == "__main__":
    unittest.main()
